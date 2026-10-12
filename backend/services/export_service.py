"""Export service for LegalEase.

Generates professional documents in TXT, DOCX (python-docx), and PDF (ReportLab) formats.
Ensures that the latest edited text is exported with complete structural fidelity.
"""

import io
import re
from typing import List, Optional

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    HRFlowable,
    KeepTogether,
)
from reportlab.pdfgen import canvas

from backend.config import settings
from backend.utils.text_utils import (
    sanitize_text,
    sanitize_filename,
    parse_document_blocks,
    parse_terms,
)


class NumberedCanvas(canvas.Canvas):
    """Canvas that computes total pages dynamically for 'Page X of Y' footers."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count: int):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))

        # Header (Top)
        self.drawString(54, letter[1] - 36, "LegalEase | AI-Powered Legal Document Generator")
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(54, letter[1] - 42, letter[0] - 54, letter[1] - 42)

        # Footer (Bottom)
        self.line(54, 45, letter[0] - 54, 45)
        self.drawString(54, 32, "Confidential Draft - Informational / Educational Use Only")
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(letter[0] - 54, 32, page_str)

        self.restoreState()


class ExportService:
    """Provides methods to export legal text drafts into TXT, DOCX, and PDF."""

    @staticmethod
    def generate_txt(
        content: str,
        document_title: str = "Legal Document",
        document_type: str = "Agreement",
    ) -> bytes:
        """Export document text into a clean UTF-8 text file."""
        clean_content = sanitize_text(content)
        separator = "=" * 70
        header = f"{separator}\n{document_title.upper()}\nDocument Type: {document_type}\n{separator}\n\n"
        footer = f"\n\n{separator}\nNOTICE:\n{settings.LEGAL_DISCLAIMER}\n{separator}\n"
        full_text = header + clean_content + footer
        return full_text.encode("utf-8")

    @staticmethod
    def generate_docx(
        content: str,
        document_title: str = "Legal Document",
        document_type: str = "Agreement",
        terms: Optional[str] = None,
    ) -> bytes:
        """Export document to Microsoft Word format with professional layout, typography, and headers/footers."""
        clean_content = sanitize_text(content)
        doc = Document()

        # Set 1-inch margins
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

            # Header
            header = section.header
            header_p = header.paragraphs[0]
            header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            h_run = header_p.add_run("LegalEase | AI-Powered Legal Document Generator")
            h_run.font.name = "Calibri"
            h_run.font.size = Pt(8.5)
            h_run.font.color.rgb = RGBColor(120, 130, 145)

            # Footer
            footer = section.footer
            footer_p = footer.paragraphs[0]
            footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            f_run = footer_p.add_run(
                "Draft Document • Generated for informational purposes only • Not formal legal advice"
            )
            f_run.font.name = "Calibri"
            f_run.font.size = Pt(8.0)
            f_run.font.italic = True
            f_run.font.color.rgb = RGBColor(140, 145, 155)

        # Base Normal Style
        normal_style = doc.styles["Normal"]
        normal_style.font.name = "Times New Roman"
        normal_style.font.size = Pt(11)
        normal_style.font.color.rgb = RGBColor(30, 30, 35)

        # Parse blocks
        blocks = parse_document_blocks(clean_content)

        # Brand Bar / Title
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(0)
        p_title.paragraph_format.space_after = Pt(12)
        r_title = p_title.add_run(document_title.upper())
        r_title.bold = True
        r_title.font.name = "Times New Roman"
        r_title.font.size = Pt(16)
        r_title.font.color.rgb = RGBColor(15, 23, 42)

        # Add optional Term Table if specific terms are present
        if terms and terms.strip():
            parsed_t = parse_terms(terms)
            if len(parsed_t) > 1:
                t_heading = doc.add_paragraph()
                t_heading.paragraph_format.space_before = Pt(8)
                t_heading.paragraph_format.space_after = Pt(4)
                r_th = t_heading.add_run("KEY TERMS & SUMMARY OF CONDITIONS")
                r_th.bold = True
                r_th.font.name = "Times New Roman"
                r_th.font.size = Pt(11)
                r_th.font.color.rgb = RGBColor(30, 41, 59)

                table = doc.add_table(rows=1, cols=2)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False

                hdr_cells = table.rows[0].cells
                hdr_cells[0].text = "Clause #"
                hdr_cells[1].text = "Agreed Condition / Term"

                # Style header row
                for idx, cell in enumerate(hdr_cells):
                    cell.width = Inches(1.0 if idx == 0 else 5.2)
                    shading = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w')))
                    cell._tc.get_or_add_tcPr().append(shading)
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.bold = True
                            run.font.size = Pt(9.5)
                            run.font.name = "Times New Roman"

                for i, term in enumerate(parsed_t, 1):
                    row_cells = table.add_row().cells
                    row_cells[0].width = Inches(1.0)
                    row_cells[1].width = Inches(5.2)
                    row_cells[0].text = f"Term {i}"
                    row_cells[1].text = term
                    for cell in row_cells:
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after = Pt(2)
                            p.paragraph_format.space_before = Pt(2)
                            for run in p.runs:
                                run.font.size = Pt(9.5)
                                run.font.name = "Times New Roman"

                # Spacer after table
                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_after = Pt(8)

        # Render parsed blocks
        for block in blocks:
            b_type = block.get("type")
            text = block.get("text", "")

            # Avoid re-adding the main title if already added
            if b_type == "title" and text.lower() == document_title.lower():
                continue

            if b_type in ("title", "heading_1"):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(text)
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(13)
                run.font.color.rgb = RGBColor(30, 41, 59)

            elif b_type == "heading_2":
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(text)
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(11.5)
                run.font.color.rgb = RGBColor(51, 65, 85)

            elif b_type == "bullet":
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.15
                run = p.add_run(text)
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)

            elif b_type == "numbered":
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.15
                prefix = block.get("prefix", "")
                r_p = p.add_run(f"{prefix} ")
                r_p.bold = True
                r_p.font.name = "Times New Roman"
                r_p.font.size = Pt(11)
                r_text = p.add_run(text)
                r_text.font.name = "Times New Roman"
                r_text.font.size = Pt(11)

            elif b_type == "signature":
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(text)
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)

            else:
                # Regular paragraph
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                run = p.add_run(text)
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)

        # Disclaimer Callout at the bottom
        p_disc = doc.add_paragraph()
        p_disc.paragraph_format.space_before = Pt(16)
        p_disc.paragraph_format.space_after = Pt(6)
        r_d = p_disc.add_run(f"LEGAL NOTICE: {settings.LEGAL_DISCLAIMER}")
        r_d.font.name = "Times New Roman"
        r_d.font.size = Pt(8.5)
        r_d.font.italic = True
        r_d.font.color.rgb = RGBColor(100, 116, 139)

        target_stream = io.BytesIO()
        doc.save(target_stream)
        return target_stream.getvalue()

    @staticmethod
    def generate_pdf(
        content: str,
        document_title: str = "Legal Document",
        document_type: str = "Agreement",
    ) -> bytes:
        """Export document to PDF using ReportLab with multi-page handling, professional typography, and footers."""
        clean_content = sanitize_text(content)
        buffer = io.BytesIO()

        # Margins: 54pt = 0.75 in
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=54,
            rightMargin=54,
            topMargin=54,
            bottomMargin=54,
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=16,
            leading=20,
            alignment=1,  # Center
            textColor=colors.HexColor("#0F172A"),
            spaceAfter=14,
        )

        h1_style = ParagraphStyle(
            "DocH1",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=16,
            textColor=colors.HexColor("#1E293B"),
            spaceBefore=10,
            spaceAfter=4,
            keepWithNext=True,
        )

        h2_style = ParagraphStyle(
            "DocH2",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=10.5,
            leading=14,
            textColor=colors.HexColor("#334155"),
            spaceBefore=8,
            spaceAfter=3,
            keepWithNext=True,
        )

        body_style = ParagraphStyle(
            "DocBody",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#1E293B"),
            alignment=4,  # Justified
            spaceAfter=6,
        )

        bullet_style = ParagraphStyle(
            "DocBullet",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=10,
            leading=13,
            leftIndent=15,
            firstLineIndent=-10,
            textColor=colors.HexColor("#1E293B"),
            spaceAfter=4,
        )

        signature_style = ParagraphStyle(
            "DocSignature",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#0F172A"),
            spaceAfter=4,
            keepWithNext=True,
        )

        disclaimer_style = ParagraphStyle(
            "DocDisclaimer",
            parent=styles["Normal"],
            fontName="Helvetica-Oblique",
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#64748B"),
            alignment=1,
        )

        story = []

        # Document Title
        story.append(Paragraph(document_title.upper(), title_style))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#CBD5E1"), spaceAfter=10))

        # Parse blocks
        blocks = parse_document_blocks(clean_content)

        for block in blocks:
            b_type = block.get("type")
            text = block.get("text", "")

            # Escape XML entities for ReportLab Paragraph
            safe_text = (
                text.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )

            # Avoid re-adding main title if duplicate
            if b_type == "title" and text.lower() == document_title.lower():
                continue

            if b_type in ("title", "heading_1"):
                story.append(Paragraph(f"<b>{safe_text}</b>", h1_style))
            elif b_type == "heading_2":
                story.append(Paragraph(f"<b>{safe_text}</b>", h2_style))
            elif b_type == "bullet":
                story.append(Paragraph(f"&bull;  {safe_text}", bullet_style))
            elif b_type == "numbered":
                prefix = block.get("prefix", "")
                story.append(Paragraph(f"<b>{prefix}</b> {safe_text}", body_style))
            elif b_type == "signature":
                story.append(Paragraph(safe_text, signature_style))
            else:
                story.append(Paragraph(safe_text, body_style))

        # Disclaimer Section at the bottom
        story.append(Spacer(1, 14))
        story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#E2E8F0"), spaceAfter=8))
        safe_disclaimer = settings.LEGAL_DISCLAIMER.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(
            Paragraph(
                f"<b>LEGAL NOTICE:</b> {safe_disclaimer}",
                disclaimer_style,
            )
        )

        # Build document with NumberedCanvas for header/footer
        doc.build(story, canvasmaker=NumberedCanvas)
        return buffer.getvalue()
